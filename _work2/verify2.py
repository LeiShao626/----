# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn
import io
p = r"D:\论文\zh_cn\潜思维链的能力与根本局限_中文翻译.docx"
d = Document(p)
print("paragraphs:", len(d.paragraphs))
inline = d.element.body.findall('.//'+qn('w:drawing'))
print("drawings(images):", len(inline))
cn=0; other=0
for para in d.paragraphs:
    for run in para.runs:
        rPr = run._element.find(qn('w:rPr'))
        if rPr is not None:
            rf = rPr.find(qn('w:rFonts'))
            if rf is not None and rf.get(qn('w:eastAsia'))=='宋体':
                cn+=1
            else:
                other+=1
        else:
            other+=1
print("runs eastAsia=宋体:", cn, "other:", other)
print("sections:", len(d.sections))
texts=[pp.text for pp in d.paragraphs if pp.text.strip()]
with io.open(r"D:\论文\_work2\dump2.txt","w",encoding="utf-8") as f:
    f.write("\n".join(texts))
print("first: ", texts[0])
