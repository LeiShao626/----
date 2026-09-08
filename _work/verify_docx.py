# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn
import os
p = r"D:\论文\zh_cn\ELF_嵌入式语言流_中文翻译.docx"
d = Document(p)
print("paragraphs:", len(d.paragraphs))
# count images in document body
from docx.oxml.ns import nsmap
inline = d.element.body.findall('.//'+qn('w:drawing'))
print("drawings:", len(inline))
# check fonts applied on a sample of runs
cn_ok = 0; other = 0; samples=[]
for para in d.paragraphs:
    for run in para.runs:
        rPr = run._element.find(qn('w:rPr'))
        if rPr is not None:
            rf = rPr.find(qn('w:rFonts'))
            if rf is not None:
                ea = rf.get(qn('w:eastAsia'))
                if ea == '宋体':
                    cn_ok += 1
                else:
                    other += 1
            else:
                other += 1
        else:
            other += 1
print("runs with eastAsia=宋体:", cn_ok, " other:", other)
# section count
print("sections:", len(d.sections))
# titles present?
texts = [pp.text for pp in d.paragraphs if pp.text.strip()]
print("\nFirst 12 non-empty paragraphs:")
for t in texts[:12]:
    print("  -", t[:60])
